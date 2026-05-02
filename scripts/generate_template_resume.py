#!/usr/bin/env python3
"""
generate_template_resume.py — Build references/templates/base-resume-template.docx

Mirrors the section structure of Milan's base resume (PROFESSIONAL SUMMARY,
PROFESSIONAL EXPERIENCE, SKILLS, EDUCATION) but fills it with clearly-template
placeholder text in [BRACKETS] so a public user can open the file and know
exactly what to replace.

The generated docx is intentionally NOT pretty. Resume formatting (fonts, table
layouts, columns) is candidate-specific. The user provides their own polish.
The template is a starting structure, not a finished design.

Usage:
    python3 scripts/generate_template_resume.py
    # writes references/templates/base-resume-template.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def main():
    repo_root = Path(__file__).resolve().parent.parent
    out = repo_root / "references" / "templates" / "base-resume-template.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Tighten default margins a bit
    for s in doc.sections:
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)

    # ── Header (name + contact line) ──────────────────────────────────────
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.add_run("[YOUR FULL NAME]")
    r.bold = True
    r.font.size = Pt(16)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run(
        "[+1 (xxx) xxx-xxxx]   |   [your.email@example.com]   |   "
        "[City, ST]   |   linkedin.com/in/[your-handle]"
    ).font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ── PROFESSIONAL SUMMARY ──────────────────────────────────────────────
    h = doc.add_paragraph()
    rh = h.add_run("PROFESSIONAL SUMMARY")
    rh.bold = True
    rh.font.size = Pt(11)

    summary = doc.add_paragraph()
    summary.add_run(
        "[Two to three sentences. Lead with years of experience and the kind of "
        "work you're known for. Mention the stack you use most. Close with what "
        "you're moving toward next. Example pattern: \"Data analyst with 6+ "
        "years across [industries], delivering [type of insights] through "
        "[tools/methods]. Now applying [emerging skill] to [outcome].\"]"
    ).font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ── PROFESSIONAL EXPERIENCE ──────────────────────────────────────────
    h = doc.add_paragraph()
    rh = h.add_run("PROFESSIONAL EXPERIENCE")
    rh.bold = True
    rh.font.size = Pt(11)

    def job(company, dates, title, location, bullets):
        line1 = doc.add_paragraph()
        rl = line1.add_run(f"{company}   |   {dates}")
        rl.bold = True
        rl.font.size = Pt(10)

        line2 = doc.add_paragraph()
        line2.add_run(f"{title}   –   {location}").font.size = Pt(10)

        for b in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(b).font.size = Pt(10)
        doc.add_paragraph()  # spacer between jobs

    job(
        "[Most Recent Company]", "[Start Mon YYYY] – [Present or End Mon YYYY]",
        "[Job Title]", "[City, ST or Remote]",
        [
            "[Action verb] [what you did] using [tool / method], achieving "
            "[quantifiable outcome — number, percent, or scope].",
            "[Action verb] [what you built or improved], reducing [metric] by "
            "[X%] / saving [Y hours per week] / unblocking [team or process].",
            "[Action verb] [collaboration or stakeholder action] for "
            "[audience], enabling [decision or outcome].",
            "[Action verb] [emerging-tool or AI / LLM application] to "
            "[accelerate, automate, or replace] [manual process].",
        ],
    )

    job(
        "[Previous Company]", "[Mon YYYY] – [Mon YYYY]",
        "[Job Title]", "[City, ST or Remote]",
        [
            "[Bullet 1 — same shape as above. Lead with the verb. Anchor in a "
            "tool or method. End in a measurable outcome.]",
            "[Bullet 2.]",
            "[Bullet 3.]",
        ],
    )

    job(
        "[Earlier Company]", "[Mon YYYY] – [Mon YYYY]",
        "[Job Title]", "[City, ST or Remote]",
        [
            "[Bullet 1.]",
            "[Bullet 2.]",
        ],
    )

    # ── SKILLS ────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    rh = h.add_run("SKILLS")
    rh.bold = True
    rh.font.size = Pt(11)

    p = doc.add_paragraph()
    rb = p.add_run("Skills: ")
    rb.bold = True
    rb.font.size = Pt(10)
    p.add_run(
        "[List your primary tools and languages. Example: SQL (advanced), "
        "Python (pandas, NumPy), Snowflake, Tableau, Power BI, dbt, Git. "
        "Group by category if you prefer.]"
    ).font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ── EDUCATION ────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    rh = h.add_run("EDUCATION")
    rh.bold = True
    rh.font.size = Pt(11)

    line1 = doc.add_paragraph()
    line1.add_run(
        "[University Name]   |   [Start Mon YYYY] – [End Mon YYYY]"
    ).font.size = Pt(10)

    line2 = doc.add_paragraph()
    line2.add_run(
        "[Degree (BS, BA, MS, etc.)], [Major]   –   [City, ST]"
    ).font.size = Pt(10)

    line3 = doc.add_paragraph()
    line3.add_run(
        "Relevant Coursework: [Course 1, Course 2, Course 3]"
    ).font.size = Pt(10)

    # ── Footer note (template authorship) ────────────────────────────────
    doc.add_paragraph()
    note = doc.add_paragraph()
    rn = note.add_run(
        "TEMPLATE — replace every value in [BRACKETS]. Delete this line "
        "before submitting. See ONBOARDING.md in the project root for the "
        "full setup walkthrough."
    )
    rn.italic = True
    rn.font.size = Pt(8)

    doc.save(out)
    print(f"wrote {out}")
    print(f"  size: {out.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
